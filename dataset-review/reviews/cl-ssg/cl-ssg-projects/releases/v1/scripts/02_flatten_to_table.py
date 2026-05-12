"""
Flatten the parsed FICHA IDI JSONs into a single analytical table.

Output: one row per *etapa* (so multi-etapa PDFs contribute 2 rows, with
``multi_etapa_duplicate`` flagging the second copy of a byte-identical pair).
Columns are organised around the 11 replicator-oriented categories documented
in ``releases/v1/docs/field_reference_for_replication.md``.

Usage
-----
    python scripts/02_flatten_to_table.py
    python scripts/02_flatten_to_table.py --json-dir releases/v1/corpus/parsed
    python scripts/02_flatten_to_table.py --out-csv releases/v1/data/derived/ficha_idi_table.csv \
                                          --out-xlsx releases/v1/data/derived/ficha_idi_table.xlsx
"""

from __future__ import annotations

import argparse
import csv
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
DERIVED = DATA / "derived"
CORPUS = ROOT / "corpus"
DEFAULT_JSON_DIR = CORPUS / "parsed"
DEFAULT_OUT_DIR = DERIVED


# ---------------------------------------------------------------------------
# Column definitions
# ---------------------------------------------------------------------------

@dataclass
class Col:
    name: str
    source: str
    type: str
    description: str


COLUMNS: list[Col] = [
    # ---------- IDENTITY ----------
    Col("codigo_bip", "iniciativa.codigo_bip", "string", "BIP project ID, format '<numeric>-<suffix>' (e.g. '40060934-0'). Same numeric ID across years means the same project."),
    Col("nombre", "iniciativa.nombre", "string", "Project name (Spanish, usually UPPERCASE)."),
    Col("tipologia", "iniciativa.tipologia", "string", "PROYECTO (works/infrastructure) | PROGRAMA (transfer/programme) | ESTUDIO BÁSICO (study)."),
    Col("descriptor", "iniciativa.descriptor", "string", "Sub-programme code (e.g. PRU, SUBTÍTULO 33)."),
    Col("etapa_actual", "clasificacion.etapa_actual", "string", "Lifecycle stage: PERFIL | PREFACTIBILIDAD | FACTIBILIDAD | DISEÑO | EJECUCION."),
    Col("proceso_presupuestario", "header.proceso_presupuestario", "int", "Budget year the FICHA was filed for."),
    Col("postula_a", "header.postula_a", "string", "Stage being applied for (often = etapa_actual)."),
    Col("situacion_solicitud", "clasificacion.situacion_solicitud", "string", "NUEVA = new request, ARRASTRE = continuing multi-year project."),
    Col("fecha_creacion_solicitud", "(document).fecha_creacion_solicitud", "date", "Request creation date (ISO YYYY-MM-DD)."),
    Col("fecha_ultima_modificacion", "(document).fecha_ultima_modificacion", "date", "Last modification date."),

    # ---------- SECTOR ----------
    Col("sector", "clasificacion.sector", "string", "High-level sector (e.g. RECURSOS HÍDRICOS)."),
    Col("subsector", "clasificacion.subsector", "string", "Sub-sector (e.g. AGUAS LLUVIAS)."),
    Col("componente_analisis", "clasificacion.componente_analisis", "string", "Analysis level: NACIONAL | REGIONAL | INTERREGIONAL | MULTIREGIONAL."),
    Col("seia", "clasificacion.seia", "string", "Environmental impact assessment status."),
    Col("area_desarrollo_indigena", "clasificacion.area_desarrollo_indigena", "bool", "True if inside an Indigenous Development Area."),

    # ---------- LOCATION ----------
    Col("ubicacion_tipo", "ubicacion.tipo", "string", "Location level: REGION | COMUNA | PROVINCIA | INTERREGIONAL | MULTIREGIONAL | NACIONAL."),
    Col("ubicacion_nombre", "ubicacion.nombre", "string", "Specific location name (e.g. VALPARAISO)."),
    Col("region", "(derived)", "string", "Region name when ubicacion_tipo=REGION."),
    Col("comuna", "(derived)", "string", "Comuna name when ubicacion_tipo=COMUNA."),
    Col("provincia", "(derived)", "string", "Provincia name when ubicacion_tipo=PROVINCIA."),
    Col("ubicacion_raw", "ubicacion.raw", "string", "Verbatim location string from the FICHA."),
    Col("distrito", "ubicacion.distrito", "string", "Electoral district number."),
    Col("circunscripcion", "ubicacion.circunscripcion", "string", "Senate constituency."),
    Col("proyecto_relacionado_codigo", "vinculacion.proyecto_relacionado.codigo", "string", "Related project's BIP id, when applicable."),
    Col("proyecto_relacionado_tipo", "vinculacion.proyecto_relacionado.tipo_relacion", "string", "Type of relation: COMPLEMENTARIO | SUSTITUTO | DEPENDE_DE."),

    # ---------- WHY ----------
    Col("justificacion", "narrative.justificacion", "text", "Spanish narrative explaining the problem (often the climate driver)."),

    # ---------- WHAT ----------
    Col("descripcion_etapa", "narrative.descripcion_etapa", "text", "Spanish narrative describing planned activities and methodology."),
    Col("proposito", "resumen_resultados.indicadores.proposito", "text", "Project purpose / strategic objective."),
    Col("componentes_n", "(derived)", "int", "Count of explicit components."),
    Col("componentes_text", "resumen_resultados.indicadores.componentes", "text", "Components joined by ' | '."),
    Col("indicadores_proposito_text", "resumen_resultados.indicadores.indicadores_proposito", "text", "Purpose indicators joined by ' | '."),
    Col("indicadores_componentes_text", "resumen_resultados.indicadores.indicadores_componentes", "text", "Component indicators joined by ' | '."),

    # ---------- BENEFICIARIES ----------
    Col("beneficiarios_hombres", "resumen_resultados.beneficiarios_directos.hombres", "int", "Direct beneficiaries — male (UNIT VARIES — read descripcion_etapa for unit)."),
    Col("beneficiarios_mujeres", "resumen_resultados.beneficiarios_directos.mujeres", "int", "Direct beneficiaries — female (same caveat)."),
    Col("beneficiarios_total", "resumen_resultados.beneficiarios_directos.total", "int", "Total direct beneficiaries (same caveat)."),

    # ---------- COST ----------
    Col("costo_total_M_CLP", "solicitud_financiamiento.totales.costo_total", "int", "Total cost in thousands of CLP (M$)."),
    Col("costo_total_CLP", "(derived)", "int", "Total cost in CLP (= costo_total_M_CLP × 1000)."),
    Col("costo_total_USD", "(derived)", "float", "Total cost in USD using tipo_cambio_clp_per_usd."),
    Col("tipo_cambio_clp_per_usd", "solicitud_financiamiento.tipo_cambio_clp_per_usd", "float", "Reference CLP/USD exchange rate."),
    Col("moneda_presupuesto_year", "solicitud_financiamiento.moneda_presupuesto.year", "int", "Reference budget year for costs."),
    Col("moneda_presupuesto_factor", "solicitud_financiamiento.moneda_presupuesto.factor", "float", "Inflation factor (DIPRES standard)."),
    Col("fuentes_financiamiento", "(derived)", "string", "Pipe-separated unique funding sources (e.g. 'F.N.D.R. | MUNICIPAL')."),
    Col("n_fuentes", "(derived)", "int", "Number of distinct funding sources."),
    Col("pagado_al_inicio_M", "solicitud_financiamiento.totales.pagado_al_inicio_periodo", "int", "Amount already paid at the start of the budget period (M$)."),
    Col("solicitado_periodo_actual_M", "solicitud_financiamiento.totales.solicitado_periodo_actual", "int", "Amount requested for the budget year (M$)."),
    Col("solicitado_periodos_siguientes_M", "solicitud_financiamiento.totales.solicitado_periodos_siguientes", "int", "Amount requested for subsequent years (M$)."),
    Col("aportes_directos_count", "(derived)", "int", "Number of work items in programación de la inversión."),
    Col("otros_aportes_count", "(derived)", "int", "Number of in-kind / partner contributions."),
    Col("otros_aportes_total_M", "(derived)", "int", "Sum of otros_aportes costo_total_etapa_programada (M$)."),

    # ---------- DURATION ----------
    Col("duracion_meses", "resumen_resultados.duracion_meses", "int", "Total etapa duration in months (from resumen)."),
    Col("aportes_directos_max_meses", "(derived)", "int", "Longest item duration across aportes_directos (months)."),

    # ---------- IMPLEMENTATION ----------
    Col("institucion_formuladora", "instituciones_participantes.institucion_formuladora", "string", "Lead implementing agency."),
    Col("instituciones_financieras", "(derived)", "string", "Pipe-separated list of financing institutions."),
    Col("instituciones_tecnicas", "(derived)", "string", "Pipe-separated list of technical institutions."),
    Col("funcionario_nombre", "funcionario_responsable.nombre", "string", "Name of the responsible officer."),
    Col("funcionario_institucion", "funcionario_responsable.institucion", "string", "Officer's institution."),
    Col("funcionario_cargo", "funcionario_responsable.cargo", "string", "Officer's role."),
    Col("funcionario_fono", "funcionario_responsable.fono", "string", "Contact phone."),
    Col("funcionario_correo", "funcionario_responsable.correo_electronico", "string", "Contact email."),

    # ---------- APPROVAL ----------
    Col("rate_resultado", "iniciativa.rate.resultado", "string", "RATE code: RS (approved) | FI (info needed) | OT (objected) | RE (reevaluación) | IN (incumple)."),
    Col("rate_fecha", "iniciativa.rate.fecha", "date", "RATE issue date."),
    Col("rate_institucion", "iniciativa.rate.institucion", "string", "Reviewing institution."),
    Col("admisibilidad", "header.admisibilidad", "string", "Admissibility flag from the header (Si | No)."),
    Col("fecha_postulacion_sni", "header.fecha_postulacion_sni", "date", "When the FICHA entered SNI."),
    Col("fecha_ingreso_sni", "header.fecha_ingreso_sni", "date", "When the FICHA was admitted to SNI."),
    Col("conclusiones_analisis", "conclusiones_analisis", "text", "Reviewer's narrative conclusions (often empty)."),
    Col("resultado_analisis_n", "(derived)", "int", "Number of review records in the FICHA."),
    Col("resultado_analisis_latest_rate", "(derived)", "string", "Most recent RATE code in resultado_analisis[]."),
    Col("resultado_analisis_latest_fecha", "(derived)", "date", "Date of the most recent review."),
    Col("admisibilidad_admitido", "observaciones_admisibilidad.admitido", "bool", "True for OBSERVACIONES DE ADMISIBILIDAD; false for NO ADMISIBILIDAD."),
    Col("admisibilidad_observaciones_n", "(derived)", "int", "Number of admissibility observations."),
    Col("admisibilidad_observaciones_text", "(derived)", "text", "Numbered observations joined by ' | '."),

    # ---------- EXECUTION ----------
    Col("historial_solicitudes_n", "(derived)", "int", "Number of historical funding-request rows."),
    Col("historial_ejecucion_n", "(derived)", "int", "Number of historical execution rows."),
    Col("historial_ejecucion_monto_vigente_total_M", "(derived)", "int", "Sum of monto_vigente across execution rows (M$)."),
    Col("historial_ejecucion_gasto_total_M", "(derived)", "int", "Sum of gasto_total across execution rows (M$)."),

    # ---------- SCHEMA / META ----------
    Col("etapa_index", "(etapa_index)", "int", "0 for first etapa, 1 for second copy in multi-etapa PDFs."),
    Col("multi_etapa_duplicate", "(derived)", "bool", "True if this is the second copy of an identical multi-etapa pair (filter out for primary analysis)."),
    Col("georreferenciacion_section_present", "georreferenciacion_section_present", "bool", "Whether section 13 was present in the printed FICHA. Coordinates are NOT in this dataset."),
    Col("template_variant", "(document).template_variant", "string", "PROYECTO | PROGRAMA | STUB | ESTUDIO_BASICO."),
    Col("template_version", "(document).template_version", "string", "Footer 'Versión X.YYYY' string."),
    Col("n_etapas_in_pdf", "(document).n_etapas_in_pdf", "int", "Total etapas present in the source PDF."),

    # ---------- PROVENANCE ----------
    Col("source_pdf_filename", "(extraction).source_pdf.filename", "string", ""),
    Col("source_pdf_size_bytes", "(extraction).source_pdf.size_bytes", "int", ""),
    Col("source_pdf_n_pages", "(extraction).source_pdf.n_pages", "int", ""),
    Col("source_pdf_sha256", "(extraction).source_pdf.sha256", "string", "Hash for reproducibility."),
    Col("parser_version", "(extraction).parser_version", "string", ""),
    Col("parsed_at", "(extraction).parsed_at", "datetime", ""),
]


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def get(d: dict | None, *keys: str) -> Any:
    cur: Any = d
    for k in keys:
        if cur is None:
            return None
        if isinstance(cur, dict):
            cur = cur.get(k)
        else:
            return None
    return cur


def joined(items: list[str] | None, sep: str = " | ") -> str | None:
    if not items:
        return None
    return sep.join(s for s in items if s)


def _ubic_part(etapa: dict, want: str) -> str | None:
    if get(etapa, "ubicacion", "tipo") == want:
        return get(etapa, "ubicacion", "nombre")
    return None


def _max_or_none(xs: list) -> Any:
    xs = [x for x in xs if x is not None]
    return max(xs) if xs else None


def _sum_or_none(xs: list, key: str) -> int | None:
    vals = [x.get(key) for x in xs if isinstance(x, dict) and x.get(key) is not None]
    return sum(vals) if vals else None


def _latest(xs: list, date_key: str = "fecha", code_key: str = "rate") -> tuple:
    """Return (most-recent code, most-recent date) by ISO date."""
    if not xs:
        return None, None
    dated = [x for x in xs if isinstance(x, dict) and x.get(date_key)]
    if not dated:
        return None, None
    latest = max(dated, key=lambda x: x[date_key])
    return latest.get(code_key), latest.get(date_key)


# ---------------------------------------------------------------------------
# Per-record flattener
# ---------------------------------------------------------------------------

def flatten_record(rec: dict) -> list[dict]:
    """One JSON record may produce 0 (stub), 1 (single-etapa), or 2 (multi-etapa) rows."""
    rows = []
    document = rec.get("document") or {}
    extraction = rec.get("extraction") or {}
    warnings = {w.get("code") for w in extraction.get("warnings", [])}
    is_multi_dup = "multi_etapa_identical_blocks" in warnings

    for i, etapa in enumerate(rec.get("etapas") or []):
        ini = etapa.get("iniciativa") or {}
        clas = etapa.get("clasificacion") or {}
        ubic = etapa.get("ubicacion") or {}
        narr = etapa.get("narrative") or {}
        sf = etapa.get("solicitud_financiamiento") or {}
        prog = etapa.get("programacion_inversion") or {}
        rr = etapa.get("resumen_resultados") or {}
        ind = (rr.get("indicadores") or {}) if rr else {}
        ben = (rr.get("beneficiarios_directos") or {}) if rr else {}
        ip = etapa.get("instituciones_participantes") or {}
        fr = etapa.get("funcionario_responsable") or {}
        oa = etapa.get("observaciones_admisibilidad") or {}
        hp = etapa.get("historial_presupuesto") or {}

        cost_M = get(sf, "totales", "costo_total")
        cost_CLP = cost_M * 1000 if isinstance(cost_M, int) else None
        tc = sf.get("tipo_cambio_clp_per_usd") if sf else None
        cost_USD = round(cost_CLP / tc, 2) if cost_CLP and tc else None

        sources = []
        for r in (sf.get("rows") or []):
            f = r.get("fuente")
            if f and f not in sources:
                sources.append(f)

        otros = prog.get("otros_aportes") or []
        otros_total_M = None
        if otros:
            vals = [(x.get("costo_total_etapa_programada") or {}).get("M$") for x in otros]
            vals = [v for v in vals if v is not None]
            otros_total_M = sum(vals) if vals else None

        ad_meses = [a.get("duracion_meses") for a in (prog.get("aportes_directos") or [])]
        max_meses = _max_or_none(ad_meses)

        latest_rate, latest_rate_fecha = _latest(etapa.get("resultado_analisis") or [])

        admis_obs = oa.get("observaciones") if oa else None
        admis_text = joined([f"{o.get('numero')}. {o.get('texto')}" for o in (admis_obs or [])])

        ejec = (hp.get("ejecucion_presupuestaria") or [])
        ejec_mv = _sum_or_none(ejec, "monto_vigente")
        ejec_gt = _sum_or_none(ejec, "gasto_total")

        row = {
            # IDENTITY
            "codigo_bip": ini.get("codigo_bip"),
            "nombre": ini.get("nombre"),
            "tipologia": ini.get("tipologia"),
            "descriptor": ini.get("descriptor"),
            "etapa_actual": clas.get("etapa_actual"),
            "proceso_presupuestario": get(etapa, "header", "proceso_presupuestario"),
            "postula_a": get(etapa, "header", "postula_a"),
            "situacion_solicitud": clas.get("situacion_solicitud"),
            "fecha_creacion_solicitud": document.get("fecha_creacion_solicitud"),
            "fecha_ultima_modificacion": document.get("fecha_ultima_modificacion"),

            # SECTOR
            "sector": clas.get("sector"),
            "subsector": clas.get("subsector"),
            "componente_analisis": clas.get("componente_analisis"),
            "seia": clas.get("seia"),
            "area_desarrollo_indigena": clas.get("area_desarrollo_indigena"),

            # LOCATION
            "ubicacion_tipo": ubic.get("tipo"),
            "ubicacion_nombre": ubic.get("nombre"),
            "region": _ubic_part(etapa, "REGION"),
            "comuna": _ubic_part(etapa, "COMUNA"),
            "provincia": _ubic_part(etapa, "PROVINCIA"),
            "ubicacion_raw": ubic.get("raw"),
            "distrito": ubic.get("distrito"),
            "circunscripcion": ubic.get("circunscripcion"),
            "proyecto_relacionado_codigo": get(etapa, "vinculacion", "proyecto_relacionado", "codigo"),
            "proyecto_relacionado_tipo": get(etapa, "vinculacion", "proyecto_relacionado", "tipo_relacion"),

            # WHY
            "justificacion": narr.get("justificacion"),

            # WHAT
            "descripcion_etapa": narr.get("descripcion_etapa"),
            "proposito": ind.get("proposito"),
            "componentes_n": len(ind.get("componentes") or []),
            "componentes_text": joined(ind.get("componentes")),
            "indicadores_proposito_text": joined(ind.get("indicadores_proposito")),
            "indicadores_componentes_text": joined(ind.get("indicadores_componentes")),

            # BENEFICIARIES
            "beneficiarios_hombres": ben.get("hombres"),
            "beneficiarios_mujeres": ben.get("mujeres"),
            "beneficiarios_total": ben.get("total"),

            # COST
            "costo_total_M_CLP": cost_M,
            "costo_total_CLP": cost_CLP,
            "costo_total_USD": cost_USD,
            "tipo_cambio_clp_per_usd": tc,
            "moneda_presupuesto_year": get(sf, "moneda_presupuesto", "year"),
            "moneda_presupuesto_factor": get(sf, "moneda_presupuesto", "factor"),
            "fuentes_financiamiento": joined(sources),
            "n_fuentes": len(sources),
            "pagado_al_inicio_M": get(sf, "totales", "pagado_al_inicio_periodo"),
            "solicitado_periodo_actual_M": get(sf, "totales", "solicitado_periodo_actual"),
            "solicitado_periodos_siguientes_M": get(sf, "totales", "solicitado_periodos_siguientes"),
            "aportes_directos_count": len(prog.get("aportes_directos") or []),
            "otros_aportes_count": len(otros),
            "otros_aportes_total_M": otros_total_M,

            # DURATION
            "duracion_meses": rr.get("duracion_meses") if rr else None,
            "aportes_directos_max_meses": max_meses,

            # IMPLEMENTATION
            "institucion_formuladora": ip.get("institucion_formuladora"),
            "instituciones_financieras": joined(ip.get("instituciones_financieras")),
            "instituciones_tecnicas": joined(ip.get("instituciones_tecnicas")),
            "funcionario_nombre": fr.get("nombre") if fr else None,
            "funcionario_institucion": fr.get("institucion") if fr else None,
            "funcionario_cargo": fr.get("cargo") if fr else None,
            "funcionario_fono": fr.get("fono") if fr else None,
            "funcionario_correo": fr.get("correo_electronico") if fr else None,

            # APPROVAL
            "rate_resultado": get(etapa, "iniciativa", "rate", "resultado"),
            "rate_fecha": get(etapa, "iniciativa", "rate", "fecha"),
            "rate_institucion": get(etapa, "iniciativa", "rate", "institucion"),
            "admisibilidad": get(etapa, "header", "admisibilidad"),
            "fecha_postulacion_sni": get(etapa, "header", "fecha_postulacion_sni"),
            "fecha_ingreso_sni": get(etapa, "header", "fecha_ingreso_sni"),
            "conclusiones_analisis": etapa.get("conclusiones_analisis"),
            "resultado_analisis_n": len(etapa.get("resultado_analisis") or []),
            "resultado_analisis_latest_rate": latest_rate,
            "resultado_analisis_latest_fecha": latest_rate_fecha,
            "admisibilidad_admitido": oa.get("admitido") if oa else None,
            "admisibilidad_observaciones_n": len(admis_obs or []),
            "admisibilidad_observaciones_text": admis_text,

            # EXECUTION
            "historial_solicitudes_n": len(hp.get("solicitudes_financiamiento") or []),
            "historial_ejecucion_n": len(ejec),
            "historial_ejecucion_monto_vigente_total_M": ejec_mv,
            "historial_ejecucion_gasto_total_M": ejec_gt,

            # META
            "etapa_index": etapa.get("etapa_index"),
            "multi_etapa_duplicate": bool(is_multi_dup and i > 0),
            "georreferenciacion_section_present": etapa.get("georreferenciacion_section_present"),
            "template_variant": document.get("template_variant"),
            "template_version": document.get("template_version"),
            "n_etapas_in_pdf": document.get("n_etapas_in_pdf"),

            # PROVENANCE
            "source_pdf_filename": get(extraction, "source_pdf", "filename"),
            "source_pdf_size_bytes": get(extraction, "source_pdf", "size_bytes"),
            "source_pdf_n_pages": get(extraction, "source_pdf", "n_pages"),
            "source_pdf_sha256": get(extraction, "source_pdf", "sha256"),
            "parser_version": extraction.get("parser_version"),
            "parsed_at": extraction.get("parsed_at"),
        }
        rows.append(row)
    return rows


# ---------------------------------------------------------------------------
# Output writers
# ---------------------------------------------------------------------------

def write_csv(rows: list[dict], path: Path) -> None:
    cols = [c.name for c in COLUMNS]
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols, quoting=csv.QUOTE_MINIMAL)
        w.writeheader()
        for r in rows:
            w.writerow({c: ("" if r.get(c) is None else r.get(c)) for c in cols})


def write_xlsx(rows: list[dict], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "data"

    bold = Font(bold=True, name="Arial", color="FFFFFF")
    head_fill = PatternFill("solid", start_color="1F3864")
    body = Font(name="Arial")

    headers = [c.name for c in COLUMNS]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = bold
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.freeze_panes = "B2"

    for r in rows:
        ws.append([r.get(c) for c in headers])

    # Reasonable widths; long-text columns get more.
    long_cols = {"justificacion", "descripcion_etapa", "proposito",
                 "componentes_text", "indicadores_proposito_text",
                 "indicadores_componentes_text", "admisibilidad_observaciones_text",
                 "ubicacion_raw", "nombre"}
    medium_cols = {"sector_subsector_raw", "fuentes_financiamiento", "instituciones_financieras",
                   "instituciones_tecnicas", "institucion_formuladora", "funcionario_institucion",
                   "funcionario_correo", "rate_institucion", "source_pdf_filename"}
    for idx, c in enumerate(COLUMNS, start=1):
        letter = get_column_letter(idx)
        width = 70 if c.name in long_cols else (28 if c.name in medium_cols else 16)
        ws.column_dimensions[letter].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Data dictionary sheet
    ws2 = wb.create_sheet("data_dictionary")
    ws2.append(["column", "json_source", "type", "description"])
    for cell in ws2[1]:
        cell.font = bold
        cell.fill = head_fill
    ws2.freeze_panes = "A2"
    for c in COLUMNS:
        ws2.append([c.name, c.source, c.type, c.description])
    ws2.column_dimensions["A"].width = 38
    ws2.column_dimensions["B"].width = 50
    ws2.column_dimensions["C"].width = 10
    ws2.column_dimensions["D"].width = 90
    for row in ws2.iter_rows(min_row=2):
        for cell in row:
            cell.font = body
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(path)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--json-dir", type=Path, default=DEFAULT_JSON_DIR)
    ap.add_argument("--out-csv", type=Path, default=DEFAULT_OUT_DIR / "ficha_idi_table.csv")
    ap.add_argument("--out-xlsx", type=Path, default=DEFAULT_OUT_DIR / "ficha_idi_table.xlsx")
    ap.add_argument("--limit", type=int, default=None,
                    help="Only process the first N JSON files (for quick testing).")
    args = ap.parse_args()

    paths = sorted(args.json_dir.glob("*.json"))
    if args.limit:
        paths = paths[: args.limit]
    print(f"reading {len(paths)} JSONs from {args.json_dir}")

    rows: list[dict] = []
    for p in paths:
        try:
            rec = json.loads(p.read_text())
        except Exception as e:
            print(f"  skipping {p.name}: {e}")
            continue
        rows.extend(flatten_record(rec))

    print(f"writing {len(rows)} rows")

    args.out_csv.parent.mkdir(parents=True, exist_ok=True)
    write_csv(rows, args.out_csv)
    print(f"  CSV  → {args.out_csv}")

    write_xlsx(rows, args.out_xlsx)
    print(f"  XLSX → {args.out_xlsx}")


if __name__ == "__main__":
    main()
